from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, time, timedelta
from pathlib import Path
from threading import Lock
from typing import Dict, List, Tuple

from fastapi import FastAPI, File, HTTPException, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from openpyxl import Workbook, load_workbook

DATA_PATH = Path('data/storage.json')
DATA_PATH.parent.mkdir(exist_ok=True)

DEFAULT_DATA = {
    "employees": [
        {
            "employeeId": "EMP-1001",
            "firstName": "Aisha",
            "lastName": "Khan",
            "department": "Engineering",
            "designation": "Senior Software Engineer",
            "email": "aisha.khan@example.com",
            "phone": "+1-202-555-0110",
            "dateOfJoining": "2021-04-05",
            "dateOfBirth": "1991-08-14",
            "salary": 92000,
            "allowances": 450,
            "deductions": 150,
            "status": "Active",
            "leaveBalance": 14,
        },
        {
            "employeeId": "EMP-1002",
            "firstName": "Dev",
            "lastName": "Patel",
            "department": "Finance",
            "designation": "Payroll Specialist",
            "email": "dev.patel@example.com",
            "phone": "+1-202-555-0142",
            "dateOfJoining": "2019-01-14",
            "dateOfBirth": "1988-11-02",
            "salary": 68000,
            "allowances": 320,
            "deductions": 110,
            "status": "Active",
            "leaveBalance": 9,
        },
        {
            "employeeId": "EMP-1003",
            "firstName": "Lucia",
            "lastName": "Fernandez",
            "department": "Human Resources",
            "designation": "HR Business Partner",
            "email": "lucia.fernandez@example.com",
            "phone": "+1-202-555-0189",
            "dateOfJoining": "2020-07-22",
            "dateOfBirth": "1990-05-28",
            "salary": 76000,
            "allowances": 280,
            "deductions": 95,
            "status": "Active",
            "leaveBalance": 5,
        },
    ],
    "leaves": [
        {
            "id": "LV-0001",
            "employeeId": "EMP-1003",
            "type": "Annual Leave",
            "startDate": "2024-04-12",
            "endDate": "2024-04-15",
            "reason": "Family event",
            "status": "Approved",
            "createdAt": "2024-03-30T09:30:00",
            "days": 4,
        }
    ],
    "attendance": [
        {
            "employeeId": "EMP-1001",
            "date": "2024-04-01",
            "checkIn": "09:05",
            "checkOut": "17:45",
            "hoursWorked": 8.67,
        },
        {
            "employeeId": "EMP-1001",
            "date": "2024-04-02",
            "checkIn": "09:15",
            "checkOut": "17:20",
            "hoursWorked": 8.08,
        },
        {
            "employeeId": "EMP-1002",
            "date": "2024-04-01",
            "checkIn": "08:50",
            "checkOut": "17:05",
            "hoursWorked": 8.25,
        },
    ],
}


class DataStore:
    def __init__(self, path: Path):
        self.path = path
        self.lock = Lock()
        if not self.path.exists():
            self._write(DEFAULT_DATA)

    def _read(self) -> Dict:
        with self.path.open('r', encoding='utf-8') as fh:
            return json.load(fh)

    def _write(self, data: Dict) -> None:
        with self.path.open('w', encoding='utf-8') as fh:
            json.dump(data, fh, indent=2)

    def read(self) -> Dict:
        with self.lock:
            return self._read()

    def update(self, callback) -> Dict:
        with self.lock:
            data = self._read()
            callback(data)
            self._write(data)
            return data


def ensure_employee_exists(data: Dict, employee_id: str) -> Dict:
    for employee in data['employees']:
        if employee['employeeId'] == employee_id:
            return employee
    raise HTTPException(status_code=404, detail='Employee not found')


class EmployeePayload(BaseModel):
    employeeId: str = Field(..., min_length=3)
    firstName: str
    lastName: str
    department: str
    designation: str
    email: str
    phone: str
    dateOfJoining: str
    dateOfBirth: str
    salary: float
    allowances: float = 0.0
    deductions: float = 0.0
    status: str
    leaveBalance: float = 0.0


class LeavePayload(BaseModel):
    employeeId: str
    type: str
    startDate: str
    endDate: str
    reason: str


class LeaveStatusPayload(BaseModel):
    status: str


class AttendancePayload(BaseModel):
    employeeId: str
    date: str
    checkIn: str
    checkOut: str


store = DataStore(DATA_PATH)
app = FastAPI(title='NovaHRMS API', version='1.0.0')


EXCEL_DATE_EPOCH = datetime(1899, 12, 30)

HEADER_ALIASES = {
    'name': 'name',
    'employee': 'name',
    'employeename': 'name',
    'staffname': 'name',
    'employeeid': 'employeeId',
    'empid': 'employeeId',
    'employeecode': 'employeeId',
    'empcode': 'employeeId',
    'code': 'employeeId',
    'id': 'employeeId',
    'date': 'date',
    'attendancedate': 'date',
    'workdate': 'date',
    'checkin': 'clockIn',
    'clockin': 'clockIn',
    'signin': 'clockIn',
    'intime': 'clockIn',
    'checkintime': 'clockIn',
    'clockout': 'clockOut',
    'checkout': 'clockOut',
    'signout': 'clockOut',
    'outtime': 'clockOut',
    'checkouttime': 'clockOut',
    'onduty': 'onDuty',
    'offduty': 'offDuty',
    'absent': 'absent',
    'status': 'status',
}

ABSENT_MARKERS = {'true', 'yes', 'y', '1', 'absent', 'a', 'leave'}


def _normalize_text(value) -> str:
    return ' '.join(str(value or '').strip().split())


def _normalize_name(value) -> str:
    return _normalize_text(value).lower()


def _header_key(value) -> str:
    return ''.join(ch.lower() for ch in str(value or '') if ch.isalnum())


def _resolve_header(value) -> str | None:
    key = _header_key(value)
    return HEADER_ALIASES.get(key)


def _parse_excel_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if not text:
        return False
    return text in ABSENT_MARKERS


def _parse_excel_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return (EXCEL_DATE_EPOCH + timedelta(days=float(value))).date()
        except Exception:
            return None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return datetime.fromisoformat(text).date()
        except ValueError:
            pass
        if ' ' in text:
            prefix = text.split(' ')[0]
            try:
                return date.fromisoformat(prefix)
            except ValueError:
                pass
        for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _parse_excel_time(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        try:
            total_seconds = int(round(float(value) * 86400))
        except Exception:
            return None
        hours, remainder = divmod(total_seconds, 3600)
        minutes, seconds = divmod(remainder, 60)
        hours %= 24
        return time(hour=hours, minute=minutes, second=seconds)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        candidate = text.upper().replace('.', ':')
        for fmt in ('%H:%M:%S', '%H:%M', '%I:%M %p', '%I:%M%p'):
            try:
                return datetime.strptime(candidate, fmt).time().replace(microsecond=0)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(text).time()
        except ValueError:
            return None
    return None


def _format_time(value: time) -> str:
    return value.strftime('%H:%M') if value else ''


def _hours_between(work_date: date, start: time, end: time) -> float:
    start_dt = datetime.combine(work_date, start)
    end_dt = datetime.combine(work_date, end)
    if end_dt < start_dt:
        end_dt += timedelta(days=1)
    return (end_dt - start_dt).total_seconds() / 3600


def _extract_attendance_rows(workbook) -> List[Dict]:
    rows: List[Dict] = []
    for sheet in workbook.worksheets:
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            continue
        column_map: Dict[str, int] = {}
        for idx, value in enumerate(header_row):
            if value is None:
                continue
            resolved = _resolve_header(value)
            if resolved and resolved not in column_map:
                column_map[resolved] = idx
        if 'date' not in column_map or ('name' not in column_map and 'employeeId' not in column_map):
            continue

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in row):
                continue

            raw_name = ''
            if 'name' in column_map and column_map['name'] < len(row):
                name_value = row[column_map['name']]
                raw_name = _normalize_text(name_value) if name_value is not None else ''
                if _header_key(raw_name) == 'name':
                    continue
                if raw_name.lower() in {'total', 'grand total'}:
                    continue
            employee_identifier = ''
            if 'employeeId' in column_map and column_map['employeeId'] < len(row):
                emp_value = row[column_map['employeeId']]
                employee_identifier = _normalize_text(emp_value)
            date_value = row[column_map['date']] if column_map['date'] < len(row) else None
            parsed_date = _parse_excel_date(date_value)
            if not raw_name and not employee_identifier and parsed_date is None:
                continue

            record = {
                'name': raw_name,
                'employeeId': employee_identifier,
                'date': parsed_date,
                'checkIn': None,
                'checkOut': None,
                'onDuty': None,
                'offDuty': None,
                'absent': False,
            }

            if 'clockIn' in column_map and column_map['clockIn'] < len(row):
                record['checkIn'] = _parse_excel_time(row[column_map['clockIn']])
            if 'clockOut' in column_map and column_map['clockOut'] < len(row):
                record['checkOut'] = _parse_excel_time(row[column_map['clockOut']])
            if 'onDuty' in column_map and column_map['onDuty'] < len(row):
                record['onDuty'] = _parse_excel_time(row[column_map['onDuty']])
            if 'offDuty' in column_map and column_map['offDuty'] < len(row):
                record['offDuty'] = _parse_excel_time(row[column_map['offDuty']])
            absent_value = None
            if 'absent' in column_map and column_map['absent'] < len(row):
                absent_value = row[column_map['absent']]
            status_value = None
            if 'status' in column_map and column_map['status'] < len(row):
                status_value = row[column_map['status']]

            record['absent'] = _parse_excel_bool(absent_value)
            if not record['absent'] and status_value is not None:
                status_text = str(status_value).strip().lower()
                if status_text in {'a', 'absent', 'leave'}:
                    record['absent'] = True

            rows.append(record)

    return rows

app.add_middleware(
    CORSMiddleware,
    allow_origins=['*'],
    allow_credentials=True,
    allow_methods=['*'],
    allow_headers=['*'],
)

app.mount('/src', StaticFiles(directory='src'), name='src')


@app.get('/', include_in_schema=False)
def serve_index():
    return FileResponse('index.html')


@app.get('/api/dashboard')
def get_dashboard():
    data = store.read()
    employees = data['employees']
    leaves = data['leaves']
    attendance = data['attendance']
    today = date.today()
    start_window = today - timedelta(days=30)

    total_employees = len(employees)
    active_employees = sum(1 for emp in employees if emp['status'] == 'Active')
    pending_leaves = sum(1 for leave in leaves if leave['status'] == 'Pending')
    leaves_taken = sum(
        1
        for leave in leaves
        if leave['status'] == 'Approved' and date.fromisoformat(leave['startDate']) >= start_window
    )

    current_month = today.month
    current_year = today.year
    monthly_records = [
        record
        for record in attendance
        if date.fromisoformat(record['date']).month == current_month
        and date.fromisoformat(record['date']).year == current_year
    ]
    average_hours = sum(record['hoursWorked'] for record in monthly_records) / len(monthly_records) if monthly_records else 0.0

    leave_alerts = [
        {
            'employee': f"{emp['firstName']} {emp['lastName']}",
            'remaining': emp.get('leaveBalance', 0),
        }
        for emp in employees
        if emp.get('leaveBalance', 0) <= 5
    ]

    birthdays = []
    for emp in employees:
        try:
            dob = date.fromisoformat(emp['dateOfBirth'])
        except ValueError:
            continue
        next_birthday = dob.replace(year=today.year)
        if next_birthday < today:
            next_birthday = next_birthday.replace(year=today.year + 1)
        delta = (next_birthday - today).days
        if 0 <= delta <= 30:
            birthdays.append(
                {
                    'employee': f"{emp['firstName']} {emp['lastName']}",
                    'date': next_birthday.strftime('%d %b %Y'),
                }
            )

    birthdays.sort(key=lambda item: datetime.strptime(item['date'], '%d %b %Y'))

    return {
        'totals': {
            'employees': total_employees,
            'activeEmployees': active_employees,
            'pendingLeaves': pending_leaves,
            'leavesTaken': leaves_taken,
            'averageHours': average_hours,
        },
        'leaveAlerts': leave_alerts,
        'birthdays': birthdays,
    }


@app.get('/api/employees')
def list_employees():
    data = store.read()
    return sorted(data['employees'], key=lambda item: (item['department'], item['firstName']))


@app.post('/api/employees')
def create_employee(payload: EmployeePayload):
    def _create(data: Dict):
        if any(emp['employeeId'] == payload.employeeId for emp in data['employees']):
            raise HTTPException(status_code=400, detail='Employee ID already exists')
        data['employees'].append(payload.dict())

    store.update(_create)
    return {'status': 'ok'}


@app.put('/api/employees/{employee_id}')
def update_employee(employee_id: str, payload: EmployeePayload):
    if payload.employeeId != employee_id:
        raise HTTPException(status_code=400, detail='Employee ID cannot be changed')

    def _update(data: Dict):
        for idx, employee in enumerate(data['employees']):
            if employee['employeeId'] == employee_id:
                data['employees'][idx] = payload.dict()
                break
        else:
            raise HTTPException(status_code=404, detail='Employee not found')

    store.update(_update)
    return {'status': 'ok'}


@app.delete('/api/employees/{employee_id}')
def delete_employee(employee_id: str):
    def _delete(data: Dict):
        before = len(data['employees'])
        data['employees'] = [emp for emp in data['employees'] if emp['employeeId'] != employee_id]
        if before == len(data['employees']):
            raise HTTPException(status_code=404, detail='Employee not found')
        data['leaves'] = [leave for leave in data['leaves'] if leave['employeeId'] != employee_id]
        data['attendance'] = [record for record in data['attendance'] if record['employeeId'] != employee_id]

    store.update(_delete)
    return Response(status_code=204)


@app.get('/api/leaves')
def list_leaves():
    data = store.read()
    return sorted(data['leaves'], key=lambda item: item['createdAt'], reverse=True)


def calculate_days(start: str, end: str) -> int:
    start_date = date.fromisoformat(start)
    end_date = date.fromisoformat(end)
    if end_date < start_date:
        raise HTTPException(status_code=400, detail='End date must be after start date')
    return (end_date - start_date).days + 1


@app.post('/api/leaves')
def create_leave(payload: LeavePayload):
    days = calculate_days(payload.startDate, payload.endDate)

    def _create(data: Dict):
        employee = ensure_employee_exists(data, payload.employeeId)
        leave_id = f"LV-{len(data['leaves']) + 1:04d}"
        record = {
            'id': leave_id,
            'employeeId': payload.employeeId,
            'type': payload.type,
            'startDate': payload.startDate,
            'endDate': payload.endDate,
            'reason': payload.reason,
            'status': 'Pending',
            'createdAt': datetime.utcnow().isoformat(timespec='seconds'),
            'days': days,
        }
        data['leaves'].append(record)

    store.update(_create)
    return {'status': 'ok'}


@app.post('/api/leaves/{leave_id}/status')
def update_leave_status(leave_id: str, payload: LeaveStatusPayload):
    if payload.status not in {'Approved', 'Rejected'}:
        raise HTTPException(status_code=400, detail='Invalid status')

    def _update(data: Dict):
        for leave in data['leaves']:
            if leave['id'] == leave_id:
                if leave['status'] != 'Pending':
                    raise HTTPException(status_code=400, detail='Only pending leaves can be updated')
                leave['status'] = payload.status
                if payload.status == 'Approved':
                    employee = ensure_employee_exists(data, leave['employeeId'])
                    balance = employee.get('leaveBalance', 0)
                    if balance < leave['days']:
                        raise HTTPException(status_code=400, detail='Insufficient leave balance')
                    employee['leaveBalance'] = round(balance - leave['days'], 2)
                break
        else:
            raise HTTPException(status_code=404, detail='Leave not found')

    store.update(_update)
    return {'status': 'ok'}


@app.post('/api/attendance')
def record_attendance(payload: AttendancePayload):
    try:
        check_in = datetime.strptime(payload.checkIn, '%H:%M')
        check_out = datetime.strptime(payload.checkOut, '%H:%M')
    except ValueError as exc:
        raise HTTPException(status_code=400, detail='Invalid time format') from exc
    if check_out <= check_in:
        raise HTTPException(status_code=400, detail='Check-out must be after check-in')

    hours = round((check_out - check_in).seconds / 3600, 2)

    def _record(data: Dict):
        ensure_employee_exists(data, payload.employeeId)
        existing = None
        for record in data['attendance']:
            if record['employeeId'] == payload.employeeId and record['date'] == payload.date:
                existing = record
                break
        if existing:
            existing.update(
                {
                    'checkIn': payload.checkIn,
                    'checkOut': payload.checkOut,
                    'hoursWorked': hours,
                }
            )
        else:
            data['attendance'].append(
                {
                    'employeeId': payload.employeeId,
                    'date': payload.date,
                    'checkIn': payload.checkIn,
                    'checkOut': payload.checkOut,
                    'hoursWorked': hours,
                }
            )

    store.update(_record)
    return {'status': 'ok'}


@app.post('/api/attendance/import')
async def import_attendance(file: UploadFile = File(...)):
    contents = await file.read()
    if not contents:
        raise HTTPException(status_code=400, detail='Uploaded file is empty')

    try:
        workbook = load_workbook(filename=io.BytesIO(contents), data_only=True)
    except Exception as exc:  # pragma: no cover - defensive parsing
        raise HTTPException(status_code=400, detail='Unable to read Excel workbook') from exc

    rows = _extract_attendance_rows(workbook)
    if not rows:
        raise HTTPException(status_code=400, detail='No attendance data found in workbook')

    summary = {
        'created': 0,
        'updated': 0,
        'skippedMissingEmployee': 0,
        'skippedInvalid': 0,
    }
    missing_employees: set[str] = set()

    def _import(data: Dict):
        employees_by_id = {emp['employeeId'].lower(): emp for emp in data['employees']}
        employees_by_name = {
            _normalize_name(f"{emp['firstName']} {emp['lastName']}"): emp for emp in data['employees']
        }
        index = {(record['employeeId'], record['date']): record for record in data['attendance']}
        processed_in_batch: set[tuple[str, str]] = set()

        for row in rows:
            work_date = row.get('date')
            if not isinstance(work_date, date):
                summary['skippedInvalid'] += 1
                continue

            identifier = row.get('employeeId')
            employee = None
            if isinstance(identifier, str):
                identifier_key = identifier.strip().lower()
                if identifier_key:
                    employee = employees_by_id.get(identifier_key)
            elif identifier:
                employee = employees_by_id.get(str(identifier).strip().lower())

            if employee is None:
                name_key = _normalize_name(row.get('name', ''))
                if name_key:
                    employee = employees_by_name.get(name_key)

            if employee is None:
                summary['skippedMissingEmployee'] += 1
                label = row.get('employeeId') or row.get('name')
                if label:
                    missing_employees.add(str(label))
                continue

            absent = bool(row.get('absent'))
            check_in_time = row.get('checkIn') or (None if absent else row.get('onDuty'))
            check_out_time = row.get('checkOut') or (None if absent else row.get('offDuty'))

            if absent:
                payload = {'checkIn': 'Absent', 'checkOut': 'Absent', 'hoursWorked': 0.0}
            else:
                if not (check_in_time and check_out_time):
                    summary['skippedInvalid'] += 1
                    continue
                hours_worked = round(_hours_between(work_date, check_in_time, check_out_time), 2)
                payload = {
                    'checkIn': _format_time(check_in_time),
                    'checkOut': _format_time(check_out_time),
                    'hoursWorked': hours_worked,
                }

            key = (employee['employeeId'], work_date.isoformat())
            if key in processed_in_batch:
                index[key].update(payload)
                continue

            processed_in_batch.add(key)

            if key in index:
                index[key].update(payload)
                summary['updated'] += 1
            else:
                record = {
                    'employeeId': employee['employeeId'],
                    'date': work_date.isoformat(),
                    **payload,
                }
                data['attendance'].append(record)
                index[key] = record
                summary['created'] += 1

    store.update(_import)

    return {
        'status': 'ok',
        **summary,
        'missingEmployees': sorted(missing_employees),
    }


@app.get('/api/attendance/{employee_id}/monthly/{year}/{month}')
def monthly_attendance(employee_id: str, year: int, month: int):
    data = store.read()
    ensure_employee_exists(data, employee_id)
    records = [
        record
        for record in data['attendance']
        if record['employeeId'] == employee_id
        and date.fromisoformat(record['date']).year == year
        and date.fromisoformat(record['date']).month == month
    ]
    records.sort(key=lambda item: item['date'])
    return {'records': records}


@app.get('/api/payroll/{employee_id}')
def payroll_summary(employee_id: str):
    data = store.read()
    employee = ensure_employee_exists(data, employee_id)
    base = float(employee.get('salary', 0))
    allowances = float(employee.get('allowances', 0))
    deductions = float(employee.get('deductions', 0))
    net = base + allowances - deductions
    return {
        'employee': f"{employee['firstName']} {employee['lastName']}",
        'base': base,
        'allowances': allowances,
        'deductions': deductions,
        'net': net,
    }


def build_payslip(employee: Dict) -> io.BytesIO:
    base = float(employee.get('salary', 0))
    allowances = float(employee.get('allowances', 0))
    deductions = float(employee.get('deductions', 0))
    net = base + allowances - deductions

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    pdf.setTitle('NovaHRMS Payslip')
    margin = 20 * mm
    cursor = height - margin

    pdf.setFont('Helvetica-Bold', 18)
    pdf.drawString(margin, cursor, 'NovaHRMS Payslip')
    cursor -= 20

    pdf.setFont('Helvetica', 11)
    pdf.drawString(margin, cursor, f"Employee: {employee['firstName']} {employee['lastName']} ({employee['employeeId']})")
    cursor -= 16
    pdf.drawString(margin, cursor, f"Department: {employee['department']}")
    cursor -= 16
    pdf.drawString(margin, cursor, f"Designation: {employee['designation']}")
    cursor -= 24

    pdf.setFont('Helvetica-Bold', 12)
    pdf.drawString(margin, cursor, 'Earnings')
    cursor -= 16
    pdf.setFont('Helvetica', 11)
    pdf.drawString(margin, cursor, f"Base salary")
    pdf.drawRightString(width - margin, cursor, f"$ {base:,.2f}")
    cursor -= 16
    pdf.drawString(margin, cursor, 'Allowances')
    pdf.drawRightString(width - margin, cursor, f"$ {allowances:,.2f}")
    cursor -= 24

    pdf.setFont('Helvetica-Bold', 12)
    pdf.drawString(margin, cursor, 'Deductions')
    cursor -= 16
    pdf.setFont('Helvetica', 11)
    pdf.drawString(margin, cursor, 'Deductions total')
    pdf.drawRightString(width - margin, cursor, f"$ {deductions:,.2f}")
    cursor -= 24

    pdf.setFont('Helvetica-Bold', 12)
    pdf.drawString(margin, cursor, 'Net pay')
    pdf.drawRightString(width - margin, cursor, f"$ {net:,.2f}")
    cursor -= 32

    pdf.setFont('Helvetica', 10)
    pdf.drawString(margin, cursor, f"Generated on {date.today().strftime('%d %b %Y')}")
    cursor -= 16
    pdf.drawString(margin, cursor, 'This is a system generated payslip and does not require a signature.')

    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer


@app.post('/api/payroll/{employee_id}/payslip')
def generate_payslip(employee_id: str):
    data = store.read()
    employee = ensure_employee_exists(data, employee_id)
    buffer = build_payslip(employee)
    filename = f"payslip-{employee_id}.pdf"
    return StreamingResponse(buffer, media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename={filename}'})


def dataset_for_export(dataset: str) -> Tuple[List[Dict], List[str]]:
    data = store.read()
    if dataset == 'employees':
        rows = data['employees']
        headers = ['employeeId', 'firstName', 'lastName', 'department', 'designation', 'email', 'phone', 'dateOfJoining', 'dateOfBirth', 'salary', 'allowances', 'deductions', 'status', 'leaveBalance']
    elif dataset == 'leaves':
        rows = data['leaves']
        headers = ['id', 'employeeId', 'type', 'startDate', 'endDate', 'status', 'reason', 'days']
    elif dataset == 'attendance':
        rows = data['attendance']
        headers = ['employeeId', 'date', 'checkIn', 'checkOut', 'hoursWorked']
    else:
        raise HTTPException(status_code=404, detail='Dataset not found')
    return rows, headers


@app.get('/api/export/{dataset}')
def export_dataset(dataset: str, format: str = 'csv'):
    rows, headers = dataset_for_export(dataset)
    if format == 'csv':
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, '') for key in headers})
        return Response(
            content=stream.getvalue(),
            media_type='text/csv',
            headers={'Content-Disposition': f'attachment; filename={dataset}.csv'},
        )
    if format == 'xlsx':
        wb = Workbook()
        ws = wb.active
        ws.title = dataset.capitalize()
        ws.append(headers)
        for row in rows:
            ws.append([row.get(key, '') for key in headers])
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={dataset}.xlsx'},
        )
    raise HTTPException(status_code=400, detail='Unsupported format')
